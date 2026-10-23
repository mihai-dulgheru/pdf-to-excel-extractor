import os
import tempfile
import unittest

import openpyxl
import pandas as pd

from config import Constants
from modules.excel_generator import ExcelGenerator


def _row(inv, code, origin="RO", eur=100.0, weight=10, vat="SK1", date="03.03.2025"):
    return {"company": "X", "invoice_number": inv, "nc8_code": code, "origin": origin, "destination": "SK",
            "invoice_value_eur": eur, "net_weight": weight, "shipment_date": date, "exchange_rate": 5.0,
            "value_ron": eur * 5, "vat_number": vat, "delivery_location": 1759, "delivery_condition": "FCA"}


def _df(*rows):
    return pd.DataFrame(list(rows), columns=Constants.COLUMNS)


def _data_rows(path):
    ws = openpyxl.load_workbook(path).active
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    return [r for r in rows if r["Nr Factura Marfa"] is not None]


class TestNrCrt(unittest.TestCase):
    def test_rows_of_one_invoice_share_one_number(self):
        """Customer convention: Nr Crt counts invoices inside a VAT group, not rows."""
        with tempfile.TemporaryDirectory() as td:
            out = os.path.join(td, "fresh.xlsx")
            ExcelGenerator(_df(_row("1001", "94019920"), _row("1001", "87082990"), _row("1002", "94019920"))) \
                .generate_excel(out)
            self.assertEqual([r["Nr Crt"] for r in _data_rows(out)], [1, 1, 2])


class TestAppendPerOrigin(unittest.TestCase):
    def test_new_invoice_with_two_origins_is_appended_whole_in_extraction_order(self):
        with tempfile.TemporaryDirectory() as td:
            base, out = os.path.join(td, "base.xlsx"), os.path.join(td, "out.xlsx")
            ExcelGenerator(_df(_row("1001", "94019920", "RO"), _row("1002", "94019920", "RO"))).generate_excel(base)
            ExcelGenerator(_df(_row("1003", "94019920", "RO"),
                               _row("1003", "94019920", "CZ", eur=50.0, weight=5))).generate_excel(out, base)
            rows = _data_rows(out)
            self.assertEqual([(r["Nr Factura Marfa"], r["Origine"], r["Nr Crt"]) for r in rows],
                             [(1001, "RO", 1), (1002, "RO", 2), (1003, "RO", 3), (1003, "CZ", 3)])

    def test_invoice_already_in_the_workbook_is_skipped_whole_even_if_its_rows_changed(self):
        """Re-running a PDF that an older version already wrote must not add origin rows next to the old ones."""
        with tempfile.TemporaryDirectory() as td:
            base, out = os.path.join(td, "base.xlsx"), os.path.join(td, "out.xlsx")
            ExcelGenerator(_df(_row("1001", "94019920", "RO"), _row("1002", "94019920", "RO"))).generate_excel(base)
            ExcelGenerator(_df(_row("1001", "94019920", "RO", eur=60.0),
                               _row("1001", "94019920", "CZ", eur=40.0),
                               _row("1003", "87082990", "RO"))).generate_excel(out, base)
            rows = _data_rows(out)
            self.assertEqual([(r["Nr Factura Marfa"], r["Origine"], r["Val Fact Euro"]) for r in rows],
                             [(1001, "RO", 100.0), (1002, "RO", 100.0), (1003, "RO", 100.0)])


def _formula_cells(path, sheet=None):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")}


class TestAppendKeepsFormulasRight(unittest.TestCase):
    def test_customer_formulas_below_the_table_follow_the_rows_they_point_at(self):
        """Rows inserted above shift every reference, including a GRAND TOTAL the customer typed under the table."""
        with tempfile.TemporaryDirectory() as td:
            base, out = os.path.join(td, "base.xlsx"), os.path.join(td, "out.xlsx")
            ExcelGenerator(_df(_row("1001", "94019920", vat="A1"), _row("1002", "94019920", vat="A1"),
                               _row("2001", "94019920", vat="B1"))).generate_excel(base)
            # layout now: 2,3 = A1 rows; 4 = Total A1; 5 = B1 row; 6 = Total B1
            wb = openpyxl.load_workbook(base)
            ws = wb.active
            ws["R3"] = "=G2*2"                    # a data-row formula pointing at ANOTHER row
            ws["L8"] = "GRAND TOTAL"
            ws["K8"] = "=SUM(K4,K6)"              # the two Total rows
            ws["H8"] = "=SUM(H2:H6)"
            ws["R8"] = "=K8/2"
            wb.save(base)

            ExcelGenerator(_df(_row("1000", "94019920", vat="A1"))).generate_excel(out, base)
            # new row lands at 2 (invoice 1000 sorts first); totals re-created at 5 and 7; GRAND TOTAL moves to 9
            f = _formula_cells(out)
            self.assertEqual(f["K2"], "=G2*J2")
            self.assertEqual(f["K3"], "=G3*J3")   # old row 2, moved down by one
            self.assertEqual(f["R4"], "=G3*2")    # old R3: its target row 2 moved to 3
            self.assertEqual(f["K9"], "=SUM(K5,K7)")
            self.assertEqual(f["H9"], "=SUM(H2:H7)")
            self.assertEqual(f["R9"], "=K9/2")
            self.assertEqual(f["K5"], "=SUM(K2:K4)")
            self.assertEqual(f["K7"], "=SUM(K6:K6)")


class TestAppendEdgeCases(unittest.TestCase):
    def test_nothing_new_still_writes_an_unchanged_workbook_and_reports_zero(self):
        with tempfile.TemporaryDirectory() as td:
            base, out = os.path.join(td, "base.xlsx"), os.path.join(td, "out.xlsx")
            ExcelGenerator(_df(_row("1001", "94019920"))).generate_excel(base)
            gen = ExcelGenerator(_df(_row("1001", "94019920")))
            gen.generate_excel(out, base)
            self.assertTrue(os.path.exists(out))
            self.assertEqual(len(gen.appended), 0)
            self.assertEqual([r["Nr Factura Marfa"] for r in _data_rows(out)], [1001])

    def test_append_targets_the_sheet_with_the_header_even_if_another_sheet_is_active(self):
        with tempfile.TemporaryDirectory() as td:
            base, out = os.path.join(td, "base.xlsx"), os.path.join(td, "out.xlsx")
            ExcelGenerator(_df(_row("1001", "94019920"))).generate_excel(base)
            wb = openpyxl.load_workbook(base)
            errors = wb.create_sheet("ERORI")
            errors["A1"] = 1
            wb.active = wb.sheetnames.index("ERORI")
            wb.save(base)
            gen = ExcelGenerator(_df(_row("1003", "94019920")))
            gen.generate_excel(out, base)
            self.assertEqual(len(gen.appended), 1)
            wb = openpyxl.load_workbook(out)
            self.assertEqual([r[2] for r in wb["Invoices"].iter_rows(min_row=2, values_only=True) if r[2]], [1001, 1003])
            self.assertEqual(wb["ERORI"]["A1"].value, 1)

    def test_workbook_without_the_expected_header_fails_with_a_clear_message(self):
        with tempfile.TemporaryDirectory() as td:
            base, out = os.path.join(td, "base.xlsx"), os.path.join(td, "out.xlsx")
            wb = openpyxl.Workbook()
            wb.active["A1"] = "altceva"
            wb.save(base)
            with self.assertRaises(ValueError) as ctx:
                ExcelGenerator(_df(_row("1001", "94019920"))).generate_excel(out, base)
            self.assertIn("antet", str(ctx.exception))


class TestRepointEdgeCases(unittest.TestCase):
    def test_per_group_range_follows_excel_and_does_not_absorb_the_total_row_above(self):
        """Only a range starting at the table top (row 2) grows; a range over one client's rows shifts like in Excel."""
        with tempfile.TemporaryDirectory() as td:
            base = os.path.join(td, "base.xlsx")
            ExcelGenerator(_df(_row("1001", "94019920", vat="A1"), _row("1002", "94019920", vat="A1"),
                               _row("3001", "94019920", vat="C1"), _row("3002", "94019920", vat="C1"))).generate_excel(base)
            # layout: 2,3 = A1; 4 = Total A1; 5,6 = C1; 7 = Total C1
            wb = openpyxl.load_workbook(base)
            wb.active["K9"] = "=SUM(K5:K6)"  # the customer's RON total for client C1
            wb.save(base)

            out_c = os.path.join(td, "c.xlsx")
            ExcelGenerator(_df(_row("3003", "94019920", vat="C1", date="04.03.2025"))).generate_excel(out_c, base)
            self.assertEqual(_formula_cells(out_c)["K10"], "=SUM(K5:K6)")  # new C1 row at 7, Total A1 re-created at 4

            out_a = os.path.join(td, "a.xlsx")
            ExcelGenerator(_df(_row("1003", "94019920", vat="A1", date="04.03.2025"))).generate_excel(out_a, base)
            self.assertEqual(_formula_cells(out_a)["K10"], "=SUM(K6:K7)")  # A1 grew by one row above the range

    def test_text_inside_quotes_is_not_treated_as_a_cell_reference(self):
        self.assertEqual(ExcelGenerator._repoint('=COUNTIF(L2:L7,"CZ27848663")', [("insert", 2)], {}),
                         '=COUNTIF(L2:L8,"CZ27848663")')
        self.assertEqual(ExcelGenerator._repoint('=SUMIF(L:L,"Total FR00393162433",K4)', [("insert", 3)], {}),
                         '=SUMIF(L:L,"Total FR00393162433",K5)')


if __name__ == "__main__":
    unittest.main()
